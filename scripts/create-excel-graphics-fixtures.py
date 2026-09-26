"""Synthetic, formula-free Excel graphics QA. Never reads customer documents.
Run with the bundled Python: create-excel-graphics-fixtures.py /tmp/qa-directory
"""
import copy
import sys
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, AreaChart, PieChart, DoughnutChart, ScatterChart, BubbleChart, RadarChart, StockChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AbsoluteAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D, XDRPoint2D
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.packaging.manifest import mimetypes

mimetypes.add_type("image/webp", ".webp")

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)
picture = Image.new("RGB", (400, 240), "white")
draw = ImageDraw.Draw(picture)
for box, color in [((0, 0, 199, 119), "#E11D48"), ((200, 0, 399, 119), "#16A34A"), ((0, 120, 199, 239), "#2563EB"), ((200, 120, 399, 239), "#FACC15")]:
    draw.rectangle(box, fill=color)
draw.rectangle((15, 15, 385, 225), outline="black", width=4)
draw.text((28, 25), "IMAGE QA - ALL FOUR CORNERS", fill="black")
for ext in ["png", "jpg", "gif", "bmp", "webp"]:
    picture.save(out / f"colours.{ext}")

wb = Workbook()
wb.remove(wb.active)

def data_sheet(name):
    ws = wb.create_sheet(name)
    for row in [["Month", "Budget", "Actual"], ["Jan", 12, 8], ["Feb", 18, 22], ["Mar", 15, 17], ["Apr", 24, 19], ["May", 20, 28]]:
        ws.append(row)
    ws["A9"] = "Synthetic QA data; no customer information."
    for col in "ABC":
        ws.column_dimensions[col].width = 16
    for row in ws:
        for cell in row:
            cell.font = Font(name="Arial", size=10)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="163D5C")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    ws.row_dimensions[1].height = 24
    ws.page_setup.orientation = "landscape"
    return ws

def add_standard(name, cls, setup=None):
    ws = data_sheet(name)
    chart = cls()
    chart.title = f"{name} chart — saved data"
    chart.style = 10
    chart.width = 15
    chart.height = 10
    if setup:
        setup(chart)
    chart.add_data(Reference(ws, min_col=2, max_col=3 if cls not in [PieChart, DoughnutChart] else 2, min_row=1, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
    chart.legend.position = "b"
    ws.add_chart(chart, "E2")
    return ws, chart

add_standard("Column", BarChart)
add_standard("Bar", BarChart, lambda c: setattr(c, "type", "bar"))
add_standard("Stacked", BarChart, lambda c: setattr(c, "grouping", "stacked"))
add_standard("Percent", BarChart, lambda c: setattr(c, "grouping", "percentStacked"))
add_standard("Line", LineChart)
add_standard("Area", AreaChart)
for name, cls in [("Pie", PieChart), ("Doughnut", DoughnutChart)]:
    ws, chart = add_standard(name, cls)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showCatName = True
add_standard("Radar", RadarChart, lambda c: setattr(c, "type", "filled"))

for name, cls in [("Scatter", ScatterChart), ("Bubble", BubbleChart)]:
    ws = data_sheet(name)
    chart = cls()
    chart.title = f"{name} chart — numeric X axis"
    chart.width, chart.height = 15, 10
    x = Reference(ws, min_col=2, min_row=2, max_row=6)
    y = Reference(ws, min_col=3, min_row=2, max_row=6)
    series = Series(y, x, zvalues=Reference(ws, min_col=2, min_row=2, max_row=6) if cls is BubbleChart else None, title="Actual vs budget")
    if cls is ScatterChart:
        chart.scatterStyle = "marker"
        series.marker.symbol = "circle"
        series.graphicalProperties.line.noFill = True
    chart.series.append(series)
    ws.add_chart(chart, "E2")

ws = data_sheet("Stock")
ws.delete_rows(1, 9)
for row in [["Day", "Open", "High", "Low", "Close"], ["Mon", 12, 16, 10, 15], ["Tue", 15, 18, 11, 13], ["Wed", 13, 19, 12, 17], ["Thu", 17, 23, 16, 21]]:
    ws.append(row)
chart = StockChart()
chart.title, chart.width, chart.height = "Stock OHLC", 15, 10
chart.add_data(Reference(ws, min_col=2, max_col=5, min_row=1, max_row=5), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
ws.add_chart(chart, "G2")

ws, chart = add_standard("Combo", BarChart)
line = LineChart()
line.add_data(Reference(ws, min_col=3, min_row=1, max_row=6), titles_from_data=True)
line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
line.y_axis.axId = 200
line.y_axis.title = "Secondary axis"
line.y_axis.axPos = "r"
line.y_axis.crosses = "max"
chart += line

ws = wb.create_sheet("Images only")
ws.page_setup.orientation = "landscape"
for ext, position in [("png", "A1"), ("jpg", "F1"), ("gif", "A15"), ("bmp", "F15"), ("webp", "A29")]:
    img = XLImage(out / f"colours.{ext}")
    img.width, img.height = 260, 156
    ws.add_image(img, position)
# True two-cell and absolute anchors, independent of populated cell bounds.
img = XLImage(out / "colours.png")
img.anchor = TwoCellAnchor(_from=AnchorMarker(col=5, row=28, colOff=pixels_to_EMU(8)), to=AnchorMarker(col=9, row=37), editAs="twoCell")
ws.add_image(img)
img = XLImage(out / "colours.png")
img.anchor = AbsoluteAnchor(pos=XDRPoint2D(pixels_to_EMU(12), pixels_to_EMU(770)), ext=XDRPositiveSize2D(pixels_to_EMU(260), pixels_to_EMU(156)))
ws.add_image(img)

ws = data_sheet("Page break")
for r in range(10, 85):
    ws.cell(r, 1, f"Row {r:03d} must survive")
    ws.row_dimensions[r].height = 17
for row in [27, 61]:
    img = XLImage(out / "colours.png")
    img.width, img.height = 400, 240
    ws.add_image(img, f"E{row}")

source = wb["Column"]
cs = wb.create_chartsheet("Chart only")
cs.add_chart(copy.deepcopy(source._charts[0]))
path = out / "excel-graphics-supported.xlsx"
wb.save(path)
(out / "excel-graphics-supported.manifest.json").write_text(json.dumps({sheet.title: len(sheet._charts) + len(getattr(sheet, "_images", [])) for sheet in wb._sheets}, indent=2))

# Add a genuinely embedded SVG picture (openpyxl converts some raster formats
# to PNG itself), plus explicit crop/rotation/flip to one existing PNG.
with ZipFile(path) as z:
    entries = {n: z.read(n) for n in z.namelist()}
# openpyxl transcodes BMP/WebP into PNG bytes while retaining their suffix.
# Restore native fixture bytes so the browser's real decoders are exercised.
for part in entries:
    if part.startswith("xl/media/") and part.rsplit(".", 1)[-1] in {"bmp", "webp"}:
        entries[part] = (out / f"colours.{part.rsplit('.', 1)[-1]}").read_bytes()
ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing", "a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
image_drawing = next(n for n, b in entries.items() if n.startswith("xl/drawings/drawing") and n.endswith(".xml") and b.count(b"<pic>") + b.count(b"<xdr:pic>") >= 7)
root = ET.fromstring(entries[image_drawing])
pic = root.find(".//xdr:pic", ns)
sp = pic.find("xdr:spPr", ns)
xfrm = ET.SubElement(sp, "{" + ns["a"] + "}xfrm", {"rot": "1800000", "flipH": "1"})
fill = pic.find("xdr:blipFill", ns)
ET.SubElement(fill, "{" + ns["a"] + "}srcRect", {"l": "10000", "t": "5000", "r": "10000", "b": "5000"})
entries[image_drawing] = ET.tostring(root)
# Convert the final absolute-anchor picture to self-contained SVG data.
rels_path = image_drawing.rsplit("/", 1)[0] + "/_rels/" + image_drawing.rsplit("/", 1)[1] + ".rels"
rels = ET.fromstring(entries[rels_path])
svg_id = root.findall(".//a:blip", ns)[-1].attrib["{" + ns["r"] + "}embed"]
for rel in rels:
    if rel.attrib["Id"] == svg_id:
        rel.attrib["Target"] = "../media/qa-vector.svg"
entries[rels_path] = ET.tostring(rels)
entries["xl/media/qa-vector.svg"] = b'<svg xmlns="http://www.w3.org/2000/svg" width="400" height="240"><rect width="400" height="240" fill="#bae6fd"/><circle cx="80" cy="120" r="60" fill="#2563eb"/><text x="160" y="125" font-family="Arial" font-size="28">SVG PICTURE</text></svg>'
types = ET.fromstring(entries["[Content_Types].xml"])
ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default", {"Extension": "svg", "ContentType": "image/svg+xml"})
entries["[Content_Types].xml"] = ET.tostring(types)
with ZipFile(path, "w", ZIP_DEFLATED) as z:
    for n, b in entries.items():
        z.writestr(n, b)
print(path)
